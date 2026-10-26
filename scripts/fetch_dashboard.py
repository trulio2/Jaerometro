from __future__ import annotations

import hashlib
import io
import json
import os
import re
import tempfile
import time
import unicodedata
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import quote, urljoin

import certifi
import openpyxl
import requests

from history_store import HistoryStore

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "front" / "public" / "data" / "dashboard.json"
HEADERS = {"User-Agent": "Mozilla/5.0"}
TIMEOUT = 60
CANDLE_MONTHS = 96
# Enough months for 24 points of year-over-year growth on 12-month sums.
FISCAL_MONTHS = 72
NOW = datetime.now(timezone.utc)
DATA_DIR = Path(os.environ.get("JAEROMETRO_DATA_DIR", ROOT / "data"))
# Re-fetch this far behind the last stored row so revised figures are picked up.
REVISION_DAYS = 120
STORE = HistoryStore(DATA_DIR)


def request_json(url: str):
    last_error = None

    for attempt in range(5):
        try:
            response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
            response.raise_for_status()
            return response.json()
        except Exception as error:
            last_error = error
            time.sleep(2 + attempt)

    raise last_error


def bcb_series(code: int, months: int = 36):
    window_start = (NOW - timedelta(days=months * 31)).date()

    def fetch(first, last):
        start = window_start
        # A series can start a few weeks after the window opens, so allow
        # some slack before deciding the stored history is too short.
        covered = first and date.fromisoformat(first) <= window_start + timedelta(
            days=45
        )

        if covered:
            start = date.fromisoformat(last) - timedelta(days=REVISION_DAYS)

        url = (
            f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{code}/dados"
            f"?formato=json&dataInicial={start.strftime('%d/%m/%Y')}"
            f"&dataFinal={NOW.date().strftime('%d/%m/%Y')}"
        )
        payload = request_json(url)

        # BCB sometimes answers 200 with an error object instead of rows.
        if not isinstance(payload, list):
            raise ValueError(f"unexpected BCB response: {str(payload)[:120]}")

        return [
            [
                datetime.strptime(row["data"], "%d/%m/%Y").date().isoformat(),
                float(str(row["valor"]).replace(",", ".")),
            ]
            for row in payload
        ]

    rows = STORE.series(f"bcb-{code}", fetch)
    return [(date.fromisoformat(day), value) for day, value in rows]


def monthly_points(rows, limit: int = 36, scale: float = 1.0):
    by_month = {}

    for date, value in rows:
        by_month[(date.year, date.month)] = value

    selected = sorted(by_month.items())[-limit:]

    return [
        {
            "date": f"{year}-{month:02d}",
            "label": f"{month:02d}/{str(year)[-2:]}",
            "value": round(value * scale, 4),
        }
        for (year, month), value in selected
    ]


def monthly_sum_points(rows, limit: int = 36, scale: float = 1.0):
    by_month = {}

    for date, value in rows:
        key = (date.year, date.month)
        by_month[key] = by_month.get(key, 0.0) + value

    selected = sorted(by_month.items())[-limit:]

    return [
        {
            "date": f"{year}-{month:02d}",
            "label": f"{month:02d}/{str(year)[-2:]}",
            "value": round(value * scale, 4),
        }
        for (year, month), value in selected
    ]


def annual_points(rows, limit: int = 24):
    by_year = {}

    for date, value in rows:
        by_year[date.year] = value

    selected = sorted(by_year.items())[-limit:]

    return [
        {
            "date": f"{year}-01-01",
            "label": str(year),
            "value": round(value, 4),
        }
        for year, value in selected
    ]


def months_ago(count: int) -> str:
    month_index = NOW.year * 12 + NOW.month - 1 - count
    return f"{month_index // 12}-{month_index % 12 + 1:02d}"


def yahoo_monthly(symbol: str, limit: int = 36):
    def fetch(first, last):
        # Once enough history is stored, a short range refreshes the last
        # months, including the still-open current one.
        covered = first and first <= months_ago(limit - 1)
        span = "1y" if covered else "10y"
        url = (
            "https://query1.finance.yahoo.com/v8/finance/chart/"
            f"{quote(symbol, safe='')}?range={span}&interval=1mo"
        )
        payload = request_json(url)
        result = payload["chart"]["result"][0]
        timestamps = result.get("timestamp", [])
        quote_data = result["indicators"]["quote"][0]
        opens = quote_data.get("open", [])
        highs = quote_data.get("high", [])
        lows = quote_data.get("low", [])
        closes = quote_data.get("close", [])
        by_month = {}

        for timestamp, open_price, high, low, close in zip(
            timestamps, opens, highs, lows, closes
        ):
            if close is None:
                continue

            moment = datetime.fromtimestamp(timestamp, tz=timezone.utc)
            by_month[f"{moment.year}-{moment.month:02d}"] = [
                float(open_price or close),
                float(high or close),
                float(low or close),
                float(close),
            ]

        return [[month, *values] for month, values in sorted(by_month.items())]

    key = "yahoo-" + re.sub(r"[^A-Za-z0-9-]", "", symbol)
    selected = STORE.series(key, fetch)[-limit:]

    return [
        {
            "date": month,
            "label": f"{month[5:7]}/{month[2:4]}",
            "value": round(close, 4),
            "open": round(open_price, 4),
            "high": round(high, 4),
            "low": round(low, 4),
            "close": round(close, 4),
        }
        for month, open_price, high, low, close in selected
    ]


def monthly_ohlc_from_daily(rows, limit: int = 36):
    by_month = {}

    for date, value in rows:
        by_month.setdefault((date.year, date.month), []).append(value)

    selected = sorted(by_month.items())[-limit:]

    return [
        {
            "date": f"{year}-{month:02d}",
            "label": f"{month:02d}/{str(year)[-2:]}",
            "value": round(values[-1], 4),
            "open": round(values[0], 4),
            "high": round(max(values), 4),
            "low": round(min(values), 4),
            "close": round(values[-1], 4),
        }
        for (year, month), values in selected
    ]


def combine_ohlc(left, right):
    right_by_date = {point["date"]: point for point in right}
    result = []

    for point in left:
        match = right_by_date.get(point["date"])

        if not match:
            continue

        result.append(
            {
                **point,
                "value": round(point["close"] * match["close"], 4),
                "open": round(point["open"] * match["open"], 4),
                "high": round(point["high"] * match["high"], 4),
                "low": round(point["low"] * match["low"], 4),
                "close": round(point["close"] * match["close"], 4),
            }
        )

    return result


def b3_session(symbol: str):
    """Stored [session date, close, daily change %] rows for a B3 instrument.

    The intraday history gives the session date and close; the quotation
    endpoint gives the change against the previous session's close, which
    the intraday history lacks when a move happens at the open.
    """

    def fetch(first, last):
        history = request_json(
            "https://cotacao.b3.com.br/mds/api/v1/DailyFluctuationHistory/"
            f"{symbol}"
        )
        floor = history.get("TradgFlr", {})
        quotes = floor.get("scty", {}).get("lstQtn", [])
        closes = [item for item in quotes if item.get("closPric")]
        day = floor.get("date")

        if not day or not closes:
            raise ValueError(f"no {symbol} session")

        change = closes[-1].get("prcFlcn") or None

        try:
            quotation = request_json(
                "https://cotacao.b3.com.br/mds/api/v1/InstrumentQuotation/"
                f"{symbol}"
            )
            trades = quotation.get("Trad") or []

            if trades:
                change = trades[0]["scty"]["SctyQtn"].get("prcFlcn", change)
        except Exception as error:
            print(f"[b3] {symbol} quotation failed: {error}")

        return [
            [
                day,
                float(closes[-1]["closPric"]),
                None if change is None else round(float(change), 4),
            ]
        ]

    try:
        return STORE.series(f"b3-{symbol}", fetch)
    except Exception:
        return []


def focus_series(indicator: str, limit: int = 5):
    def fetch():
        filter_value = quote(f"Indicador eq '{indicator}'", safe="")
        url = (
            "https://olinda.bcb.gov.br/olinda/servico/Expectativas/versao/v1/"
            "odata/ExpectativasMercadoAnuais"
            f"?$filter={filter_value}&$orderby=Data%20desc&$top=1000"
            "&$format=json"
        )
        rows = request_json(url)["value"]
        rows = [row for row in rows if row.get("baseCalculo") == 0]

        if not rows:
            return None

        latest = max(row["Data"] for row in rows)
        rows = [row for row in rows if row["Data"] == latest]
        rows.sort(key=lambda row: int(row["DataReferencia"]))

        return [
            {
                "date": row["DataReferencia"],
                "label": row["DataReferencia"],
                "value": round(float(row["Mediana"] or row["Media"]), 4),
            }
            for row in rows[-limit:]
        ]

    key = "focus-" + re.sub(r"[^a-z0-9]+", "-", indicator.lower()).strip("-")
    return STORE.latest(key, fetch, default=[])


def janjometro_value():
    def fetch():
        url = "https://api.janjometro.com/api/v1/post/sumPublishedCosts"
        payload = request_json(url)
        value = str(payload.get("totalCost", "")).replace("\xa0", " ").strip()
        return value or None

    return STORE.latest("janjometro", fetch, default="")


# impostometro.com.br does not always send its intermediate certificate (GitHub's
# runners get the bare leaf), so its requests are verified against certifi's
# roots plus that public intermediate rather than failing or skipping checks.
INTERMEDIATE_CERTS = Path(__file__).resolve().parent / "certs" / "sectigo-rsa-dv-ca.pem"


def ca_bundle_with_intermediates() -> str:
    bundle = Path(tempfile.gettempdir()) / "jaerometro-ca-bundle.pem"
    bundle.write_text(
        Path(certifi.where()).read_text(encoding="utf-8")
        + "\n"
        + INTERMEDIATE_CERTS.read_text(encoding="utf-8"),
        encoding="utf-8",
    )
    return str(bundle)


def fetch_impostometro():
    start = f"01/01/{NOW.year}"
    end = NOW.strftime("%d/%m/%Y")
    headers = {
        **HEADERS,
        "X-Requested-With": "XMLHttpRequest",
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Referer": "https://impostometro.com.br/",
    }
    response = requests.get(
        "https://impostometro.com.br/Contador/Brasil",
        params={"dataInicial": start, "dataFinal": end},
        headers=headers,
        timeout=TIMEOUT,
        verify=ca_bundle_with_intermediates(),
    )
    response.raise_for_status()
    data = response.json()
    raw = data["Data"].split(".")[0]
    local = datetime.strptime(raw, "%Y-%m-%dT%H:%M:%S").replace(
        tzinfo=timezone(timedelta(hours=-3))
    )

    return {
        "value": float(data["Valor"]),
        "incrementPerSecond": float(data["Incremento"]),
        "updatedAt": local.astimezone(timezone.utc)
        .isoformat()
        .replace("+00:00", "Z"),
    }


def impostometro_data():
    return STORE.latest("impostometro", fetch_impostometro)


# Google News search terms per topic. Long AND-style queries return almost
# nothing recent, so each topic uses a short query with a recency filter.
NEWS_TOPICS = [
    ("fiscal", "Fiscal", "ajuste fiscal when:7d"),
    ("juros", "Juros", 'Copom OR Selic OR "juros futuros" when:7d'),
    ("cambio", "Câmbio", "dólar real when:7d"),
    ("inflacao", "Inflação", "IPCA OR inflação when:7d"),
    ("mercado", "Mercado", "Ibovespa bolsa when:7d"),
]
NEWS_PER_TOPIC = 8
NEWS_SHOWN_PER_TOPIC = 6
NEWS_ARCHIVE_DAYS = 30


def news_key(title: str) -> str:
    """Accent- and punctuation-free title, to spot the same story twice."""
    text = unicodedata.normalize("NFKD", title.lower())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def fetch_news_topic(topic: str, label: str, query: str):
    url = (
        "https://news.google.com/rss/search"
        f"?q={quote(query)}&hl=pt-BR&gl=BR&ceid=BR:pt-419"
    )
    response = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    root = ET.fromstring(response.content)
    items = []

    for item in root.findall("./channel/item"):
        title = (item.findtext("title") or "").strip()
        source_node = item.find("source")
        source = (
            source_node.text.strip()
            if source_node is not None and source_node.text
            else "Google Notícias"
        )

        # Google appends " - Source" to every title; the source has its own field.
        if title.endswith(f" - {source}"):
            title = title[: -len(source) - 3].rstrip()

        try:
            published = parsedate_to_datetime(item.findtext("pubDate"))
        except (TypeError, ValueError):
            continue

        if not title:
            continue

        published_at = (
            published.astimezone(timezone.utc)
            .isoformat()
            .replace("+00:00", "Z")
        )
        items.append(
            {
                "id": hashlib.sha1(news_key(title).encode()).hexdigest()[:12],
                "sortKey": published_at,
                "title": title,
                "source": source,
                "url": item.findtext("link") or "",
                "publishedAt": published_at,
                "topic": label,
            }
        )

    items.sort(key=lambda item: item["publishedAt"], reverse=True)
    return items[:NEWS_PER_TOPIC]


def fetch_news():
    items = []
    errors = []

    for topic, label, query in NEWS_TOPICS:
        try:
            items.extend(fetch_news_topic(topic, label, query))
        except Exception as error:
            errors.append(f"{topic}: {error}")
            print(f"[news] {topic} failed: {error}")

    if not items:
        raise ValueError("; ".join(errors) or "no news found")

    # The same story often matches several topics; keep its first topic.
    unique = {}

    for item in items:
        unique.setdefault(item["id"], item)

    return list(unique.values())


def news_items():
    cutoff = (NOW - timedelta(days=NEWS_ARCHIVE_DAYS)).isoformat()
    archive = STORE.archive(
        "news",
        fetch_news,
        keep=lambda item: item["publishedAt"] >= cutoff[:19],
    )
    per_topic = {}
    shown = []

    for item in archive:
        count = per_topic.get(item["topic"], 0)

        if count < NEWS_SHOWN_PER_TOPIC:
            per_topic[item["topic"]] = count + 1
            shown.append({key: value for key, value in item.items() if key != "sortKey"})

    return shown


def latest_revenue_report_url(year: int):
    page_url = (
        "https://www.gov.br/receitafederal/pt-br/centrais-de-conteudo/"
        "publicacoes/relatorios/arrecadacao-federal/"
        f"{year}"
    )
    response = requests.get(page_url, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    pattern = re.compile(
        r'href="([^"]*analise-mensal-([a-z]{3})-\d{4}-anexo\.xlsx)/view"',
        re.IGNORECASE,
    )
    months = {
        "jan": 1,
        "fev": 2,
        "mar": 3,
        "abr": 4,
        "mai": 5,
        "jun": 6,
        "jul": 7,
        "ago": 8,
        "set": 9,
        "out": 10,
        "nov": 11,
        "dez": 12,
    }
    matches = pattern.findall(response.text)

    if not matches:
        return None

    best_url, best_month = max(
        matches, key=lambda item: months.get(item[1].lower(), 0)
    )
    return urljoin(page_url, best_url) + "/@@download/file"


def fetch_revenue_rows():
    report_url = None

    for year in (NOW.year, NOW.year - 1):
        try:
            report_url = latest_revenue_report_url(year)
        except Exception:
            report_url = None

        if report_url:
            break

    if not report_url:
        raise ValueError("no revenue report found")

    response = requests.get(report_url, headers=HEADERS, timeout=TIMEOUT)
    response.raise_for_status()
    workbook = openpyxl.load_workbook(
        io.BytesIO(response.content), data_only=True, read_only=True
    )
    sheet_name = next(
        (name for name in workbook.sheetnames if name.startswith("Tabela III")),
        None,
    )

    if not sheet_name:
        raise ValueError("revenue table not found in report")

    sheet = workbook[sheet_name]
    month_names = {
        "JAN": 1,
        "FEV": 2,
        "MAR": 3,
        "ABR": 4,
        "MAIO": 5,
        "MAI": 5,
        "JUN": 6,
        "JUL": 7,
        "AGO": 8,
        "SET": 9,
        "OUT": 10,
        "NOV": 11,
        "DEZ": 12,
    }
    start_year = NOW.year - 5

    for row in sheet.iter_rows(min_row=1, max_row=6, values_only=True):
        text = " ".join(str(value) for value in row if value)
        match = re.search(r"JANEIRO DE (\d{4})", text, re.IGNORECASE)

        if match:
            start_year = int(match.group(1))
            break

    current_year = start_year
    values = []

    for row in sheet.iter_rows(min_row=8, values_only=True):
        first = str(row[0]).strip().upper() if row and row[0] else ""

        if first.startswith("JAN-DEZ"):
            years = re.findall(r"\d{4}", first)

            if years:
                current_year = int(years[-1]) + 1

            continue

        month = month_names.get(first)
        total = row[9] if len(row) > 9 else None

        if month and isinstance(total, (int, float)):
            values.append([f"{current_year}-{month:02d}", float(total)])

    if not values:
        raise ValueError("revenue table has no monthly totals")

    return values


def revenue_points(limit: int = 36):
    rows = STORE.series("rfb-revenue", lambda first, last: fetch_revenue_rows())

    return [
        {
            "date": month,
            "label": f"{month[5:7]}/{month[2:4]}",
            "value": round(value / 1000, 4),
        }
        for month, value in rows[-limit:]
    ]


def compact_brl(text: str) -> str:
    """'R$ 117.595.594,73' -> 'R$ 117,6 mi', to fit a small tile."""
    digits = re.sub(r"[^\d,]", "", text).replace(",", ".")

    try:
        value = float(digits)
    except ValueError:
        return text

    for size, unit in ((1e12, "tri"), (1e9, "bi"), (1e6, "mi"), (1e3, "mil")):
        if value >= size:
            return f"R$ {number(value / size, 1)} {unit}"

    return f"R$ {number(value)}"


def number(value: float, decimals: int = 2):
    formatted = f"{value:,.{decimals}f}"
    return formatted.replace(",", "X").replace(".", ",").replace("X", ".")


def signed(value: float, decimals: int = 2, suffix: str = ""):
    signal = "+" if value > 0 else ""
    return f"{signal}{number(value, decimals)}{suffix}"


def change(current: float, previous: float):
    if not previous:
        return 0.0
    return (current / previous - 1) * 100


def diff(current: float, previous: float):
    return current - previous


def latest(points):
    return points[-1]["value"] if points else 0.0


def previous(points):
    return points[-2]["value"] if len(points) > 1 else latest(points)


def trend(value: float):
    if value > 0:
        return "up"
    if value < 0:
        return "down"
    return "flat"


def series_map(points):
    return {point["date"]: point["value"] for point in points}


def as_of(points):
    if points:
        return points[-1]["date"]

    return NOW.isoformat().replace("+00:00", "Z")


def severity_for(value: float):
    if value < 0:
        return "critical"
    return "normal"


def chart_point_severity(points, rule):
    return [
        {**point, "severity": rule(point["value"])}
        for point in points
    ]


def rolling_sum(points, window: int = 12, limit: int = 24):
    result = []

    for index in range(window - 1, len(points)):
        chunk = points[index - window + 1 : index + 1]
        result.append(
            {
                **points[index],
                "value": round(sum(point["value"] for point in chunk), 4),
            }
        )

    return result[-limit:]


def yoy_growth(points, limit: int = 24):
    """Percent change of each point against the same month a year before."""
    by_date = {point["date"]: point["value"] for point in points}
    result = []

    for point in points:
        year, month = point["date"][:4], point["date"][5:7]
        before = by_date.get(f"{int(year) - 1}-{month}")

        if before:
            result.append(
                {**point, "value": round((point["value"] / before - 1) * 100, 2)}
            )

    return result[-limit:]


def ratio_points(numerators, denominators, limit: int = 24):
    """numerator / denominator in percent, for the months both series have."""
    by_date = {point["date"]: point["value"] for point in denominators}

    return [
        {**point, "value": round(point["value"] / by_date[point["date"]] * 100, 2)}
        for point in numerators
        if by_date.get(point["date"])
    ][-limit:]


def rolling_average(points, window: int = 12, limit: int = 24):
    result = []

    for index in range(window - 1, len(points)):
        chunk = points[index - window + 1 : index + 1]
        result.append(
            {
                **points[index],
                "value": round(
                    sum(point["value"] for point in chunk) / window, 4
                ),
            }
        )

    return result[-limit:]


def rolling_compound(points, window: int = 12, limit: int = 24):
    result = []

    for index in range(window - 1, len(points)):
        chunk = points[index - window + 1 : index + 1]
        factor = 1.0

        for point in chunk:
            factor *= 1 + point["value"] / 100

        result.append(
            {
                **points[index],
                "value": round((factor - 1) * 100, 4),
            }
        )

    return result[-limit:]


# Daily market moves that signal stress. Thresholds are (warning, critical).
DI_JUMP_BPS = (50, 100)
USD_JUMP_PCT = (2, 3)
IBOV_DROP_PCT = (-3, -5)
# A session older than this no longer counts as "today's" move.
SESSION_MAX_AGE_DAYS = 4

SEVERITY_ORDER = {"critical": 0, "warning": 1, "info": 2}


def recent_session(day: str) -> bool:
    age = NOW.date() - date.fromisoformat(day)
    return age.days <= SESSION_MAX_AGE_DAYS


def short_date(day: str) -> str:
    return f"{day[8:10]}/{day[5:7]}"


# January DI contracts that make up the curve: the next nine years.
DI_CURVE_YEARS = range(NOW.year + 1, NOW.year + 10)
# How far back the "a month ago" curve looks, and the slack allowed when no
# session falls exactly on that day (weekends, holidays, missed runs).
CURVE_LOOKBACK_DAYS = 30
CURVE_LOOKBACK_SLACK_DAYS = 5


def di_symbol(year: int) -> str:
    return f"DI1F{year % 100:02d}"


def di_curve_series(sessions_by_year):
    """Curve points for the latest session, the one before it, and ~1 month ago."""
    latest = max(
        (rows[-1][0] for rows in sessions_by_year.values() if rows), default=None
    )

    if latest is None:
        return []

    target = date.fromisoformat(latest) - timedelta(days=CURVE_LOOKBACK_DAYS)
    today, previous, month_ago = [], [], []

    for year, rows in sorted(sessions_by_year.items()):
        point = {"date": str(year), "label": f"jan/{year % 100:02d}"}
        row = rows[-1] if rows else None

        if not row or row[0] != latest:
            continue

        today.append({**point, "value": round(row[1], 3)})

        if len(row) >= 3 and row[2] is not None:
            previous.append(
                {**point, "value": round(row[1] / (1 + row[2] / 100), 3)}
            )

        # Closest stored session on or before the target, within the slack.
        older = [
            stored
            for stored in rows
            if date.fromisoformat(stored[0]) <= target
            and (target - date.fromisoformat(stored[0])).days
            <= CURVE_LOOKBACK_SLACK_DAYS
        ]

        if older:
            month_ago.append({**point, "value": round(older[-1][1], 3)})

    series = [
        {
            "id": "di-curve-today",
            "label": f"Sessão {short_date(latest)}",
            "color": "#ff4d5e",
            "axis": "left",
            "type": "area",
            "decimals": 2,
            "suffix": "%",
            "points": today,
        },
        {
            "id": "di-curve-previous",
            "label": "Sessão anterior",
            "color": "#c9d6d7",
            "axis": "left",
            "dashed": True,
            "decimals": 2,
            "suffix": "%",
            "points": previous,
        },
    ]

    # Only once the store holds a month of history for most of the curve.
    if len(month_ago) >= len(today) / 2:
        series.append(
            {
                "id": "di-curve-month",
                "label": "1 mês atrás",
                "color": "#f5c542",
                "axis": "left",
                "dashed": True,
                "decimals": 2,
                "suffix": "%",
                "points": month_ago,
            }
        )

    return [item for item in series if item["points"]]


def di_daily_series(sessions_by_year, years, limit: int = 250):
    """Daily closes of the given January contracts, oldest first."""
    colors = ["#ff4d5e", "#f5c542"]
    series = []

    for year, color in zip(years, colors):
        rows = sessions_by_year.get(year) or []
        series.append(
            {
                "id": f"di-{year}",
                "label": f"DI jan/{year % 100:02d}",
                "color": color,
                "axis": "left",
                "decimals": 2,
                "suffix": "%",
                "points": [
                    {
                        "date": row[0],
                        "label": short_date(row[0]),
                        "value": round(row[1], 3),
                    }
                    for row in rows[-limit:]
                ],
            }
        )

    return series


def daily_market_alerts(di_sessions, ibov_sessions, usd_daily):
    alerts = []

    def add(severity, title, message, day):
        alerts.append((severity, title, message, day))

    jumps = []

    for name, rows in di_sessions.items():
        # Rows stored before the daily change was recorded have no third item.
        if not rows or len(rows[-1]) < 3 or rows[-1][2] is None:
            continue

        day, close, pct = rows[-1][:3]

        if not recent_session(day):
            continue

        before = close / (1 + pct / 100)
        jumps.append((round((close - before) * 100), name, before, close, day))

    if jumps:
        bps, name, before, close, day = max(jumps)

        if bps >= DI_JUMP_BPS[0]:
            detail = ", ".join(
                f"{item_name} {signed(item_bps, 0)} bps"
                for item_bps, item_name, *_ in sorted(jumps, reverse=True)
            )
            add(
                "critical" if bps >= DI_JUMP_BPS[1] else "warning",
                "Disparada nos Juros Futuros",
                f"{name} saltou de {number(before)}% para {number(close)}% "
                f"em {short_date(day)} ({detail}). "
                "Mercado exige prêmio maior para financiar a dívida.",
                day,
            )

    if len(usd_daily) >= 2 and recent_session(usd_daily[-1][0].isoformat()):
        (day, value), (_, before) = usd_daily[-1], usd_daily[-2]
        move = change(value, before)

        if move >= USD_JUMP_PCT[0]:
            add(
                "critical" if move >= USD_JUMP_PCT[1] else "warning",
                "Alta Forte do Dólar",
                f"Dólar PTAX subiu {number(move)}% em {short_date(day.isoformat())}, "
                f"de R$ {number(before, 4)} para R$ {number(value, 4)}.",
                day.isoformat(),
            )

    if ibov_sessions and ibov_sessions[-1][2] is not None:
        day, close, move = ibov_sessions[-1][:3]

        if recent_session(day) and move <= IBOV_DROP_PCT[0]:
            add(
                "critical" if move <= IBOV_DROP_PCT[1] else "warning",
                "Queda Forte da Bolsa",
                f"IBOVESPA caiu {number(abs(move))}% em {short_date(day)}, "
                f"fechando em {number(close, 0)} pontos.",
                day,
            )

    return alerts


def build_alerts(
    debt, real_rate, flow, ipca, revenue, money, ewz, ibov, focus_ipca, daily=()
):
    alerts = []

    def add(severity, title, message, at=None):
        alerts.append(
            {
                "id": f"alert-{len(alerts) + 1}",
                "severity": severity,
                "title": title,
                "message": message,
                "at": at or NOW.isoformat().replace("+00:00", "Z"),
            }
        )

    for severity, title, message, day in daily:
        add(severity, title, message, day)

    if debt > 80:
        add(
            "critical",
            "Dívida Pública em Zona de Risco",
            f"Dívida bruta do governo geral em {number(debt)}% do PIB. "
            "Acima da barreira de 80%, com trajetória de alta.",
        )

    if real_rate > 7:
        add(
            "warning",
            "Juro Real Elevado",
            f"Juro real ex-post em {number(real_rate)}%. "
            "Custo de carregamento da dívida pressiona o resultado fiscal.",
        )

    if flow < 0:
        add(
            "warning",
            "Saída Líquida de Capital",
            f"Fluxo financeiro externo negativo em US$ {number(abs(flow), 1)} bi "
            "no mês. Pressão sobre o câmbio.",
        )

    if ipca > 4.5:
        add(
            "warning",
            "Inflação Acima da Meta",
            f"IPCA acumulado em 12 meses em {number(ipca)}%. "
            "Risco de reprecificação da curva de juros.",
        )

    revenue_change = change(latest(revenue), previous(revenue))

    if revenue_change < -10:
        add(
            "warning",
            "Queda na Arrecadação",
            f"Arrecadação federal caiu {number(abs(revenue_change))}% "
            "na comparação mensal.",
        )

    money_change = change(latest(money), previous(money))

    if money_change > 2:
        add(
            "warning",
            "Expansão Monetária",
            f"Base monetária cresceu {number(money_change)}% no mês. "
            "Acompanhar impacto sobre inflação e câmbio.",
        )

    ewz_change = change(latest(ewz), previous(ewz))
    ibov_change = change(latest(ibov), previous(ibov))

    if ewz_change < -5 or ibov_change < -5:
        add(
            "warning",
            "Estresse no Mercado Acionário",
            f"EWZ {signed(ewz_change)}% e IBOVESPA {signed(ibov_change)}% "
            "no mês. Piora na percepção de risco.",
        )

    if latest(focus_ipca) > 4.5:
        add(
            "warning",
            "Expectativa de Inflação Elevada",
            f"Mediana Focus para o IPCA em {number(latest(focus_ipca))}%. "
            "Expectativas desancoradas da meta.",
        )

    add(
        "info",
        "Sistema Iniciado",
        "Monitoramento macro-financeiro carregado a partir da base pública.",
    )

    # Most severe first; sorted() is stable, so rule order breaks ties.
    alerts.sort(key=lambda alert: SEVERITY_ORDER[alert["severity"]])
    return alerts[:8]


def main():
    debt_points = monthly_points(bcb_series(13762))
    debt_brl_points = monthly_points(bcb_series(13761), scale=0.000001)
    usd_points = monthly_points(bcb_series(3698))
    ipca_points = monthly_points(bcb_series(433))
    ipca12_points = monthly_points(bcb_series(13522))
    selic_points = monthly_points(bcb_series(432))
    flow_points = monthly_sum_points(bcb_series(13961), scale=0.001)
    current_month = f"{NOW.year}-{NOW.month:02d}"

    if flow_points and flow_points[-1]["date"] == current_month:
        flow_points = flow_points[:-1]

    interest_points = monthly_points(
        bcb_series(4616, months=FISCAL_MONTHS), limit=FISCAL_MONTHS, scale=0.001
    )
    central_interest_points = monthly_points(
        bcb_series(4606, months=FISCAL_MONTHS), limit=FISCAL_MONTHS, scale=0.001
    )
    money_points = monthly_points(bcb_series(1788), scale=0.000001)
    pib_points = monthly_points(bcb_series(4380), scale=0.000001)
    primary_points = monthly_points(bcb_series(5793))
    nominal_points = monthly_points(bcb_series(5727))
    gdp_growth_points = annual_points(bcb_series(7326, months=24 * 12), limit=24)

    selic_by_month = series_map(selic_points)
    ipca12_by_month = series_map(ipca12_points)
    real_points = []

    for point in ipca12_points:
        selic = selic_by_month.get(point["date"])
        ipca = ipca12_by_month.get(point["date"])

        if selic is None or ipca is None:
            continue

        real = ((1 + selic / 100) / (1 + ipca / 100) - 1) * 100
        real_points.append(
            {
                "date": point["date"],
                "label": point["label"],
                "value": round(real, 4),
            }
        )

    revenue = revenue_points(limit=FISCAL_MONTHS)
    ewz_candles = yahoo_monthly("EWZ", limit=CANDLE_MONTHS)
    ibov_candles = yahoo_monthly("^BVSP", limit=CANDLE_MONTHS)
    btc_usd_candles = yahoo_monthly("BTC-USD", limit=CANDLE_MONTHS)
    usd_daily = bcb_series(1, months=CANDLE_MONTHS)
    usd_candles = monthly_ohlc_from_daily(usd_daily, limit=CANDLE_MONTHS)
    btc_brl_candles = combine_ohlc(btc_usd_candles, usd_candles)
    ewz_points = ewz_candles[-36:]
    ibov_points = ibov_candles[-36:]
    focus_ipca = focus_series("IPCA")
    focus_pib = focus_series("PIB Total")
    di_by_year = {
        year: b3_session(di_symbol(year))
        for year in sorted({*DI_CURVE_YEARS, 2032, 2035})
    }
    di_sessions = {
        "DI jan/32": di_by_year[2032],
        "DI jan/35": di_by_year[2035],
    }
    ibov_sessions = b3_session("IBOV")
    janjometro = janjometro_value()
    impostometro_live = impostometro_data()

    di_curve = di_curve_series(
        {year: di_by_year[year] for year in DI_CURVE_YEARS}
    )
    di_daily = di_daily_series(di_by_year, (2032, 2035))

    usd_rolling = rolling_average(usd_points)
    ewz_rolling = rolling_average(ewz_points)
    ibov_rolling = rolling_average(ibov_points)
    flow_rolling = rolling_sum(flow_points)
    revenue_rolling = rolling_sum(revenue, limit=FISCAL_MONTHS)
    interest_rolling = rolling_sum(interest_points, limit=FISCAL_MONTHS)
    central_interest_rolling = rolling_sum(
        central_interest_points, limit=FISCAL_MONTHS
    )
    revenue_growth = yoy_growth(revenue_rolling)
    interest_growth = yoy_growth(interest_rolling)
    interest_share = ratio_points(central_interest_rolling, revenue_rolling)

    def ipca_matching(points):
        dates = {point["date"] for point in points}
        return [point for point in ipca12_points if point["date"] in dates]

    def to_trillions(points):
        return [
            {**point, "value": round(point["value"] / 1000, 4)}
            for point in points[-24:]
        ]

    # SGS 5793 follows the NFSP convention (positive = deficit), so flip it
    # to read as a result: surplus above zero, deficit below.
    primary_12m = [
        {**point, "value": -point["value"]} for point in primary_points[-24:]
    ]
    # SGS 5727 (nominal result: primary minus interest) uses the same convention.
    nominal_12m = [
        {**point, "value": -point["value"]} for point in nominal_points[-24:]
    ]
    ipca_rolling = rolling_compound(ipca_points)
    money_rolling = rolling_average(money_points)
    real_rolling = rolling_average(real_points)
    selic_rolling = rolling_average(selic_points)
    pib_rolling = rolling_sum(pib_points)

    debt = latest(debt_points)
    usd = latest(usd_points)
    ipca = latest(ipca12_points)
    selic = latest(selic_points)
    real_rate = latest(real_points)
    flow = latest(flow_points)
    ibov = latest(ibov_points)
    ewz = latest(ewz_points)

    usd_delta = change(usd, previous(usd_points))
    ipca_delta = diff(ipca, previous(ipca12_points))
    selic_delta = diff(selic, previous(selic_points))
    real_delta = diff(real_rate, previous(real_points))
    debt_delta = diff(debt, previous(debt_points))
    ibov_delta = change(ibov, previous(ibov_points))
    ewz_delta = change(ewz, previous(ewz_points))
    flow_change = flow - previous(flow_points)
    interest_12m = latest(interest_rolling)
    interest_delta = latest(interest_growth)
    debt_brl = latest(debt_brl_points)

    metrics = [
        {
            "id": "debt",
            "label": "DÍVIDA BRUTA / PIB",
            "description": "Dívida do governo geral",
            "value": f"{number(debt)}%",
            "detail": f"R$ {number(debt_brl)} tri",
            "delta": signed(debt_delta, 2, " p.p."),
            "deltaLabel": "no mês",
            "trend": trend(debt_delta),
            "tone": "negative",
            "severity": "critical" if debt > 80 else "warning",
            "icon": "risk",
            "source": "BCB SGS 13762 / 13761",
            "updatedAt": as_of(debt_points),
        },
        {
            "id": "usd",
            "label": "DÓLAR",
            "description": "Cotação do dólar em reais",
            "value": f"R$ {number(usd)}",
            "delta": signed(usd_delta, 2, "%"),
            "deltaLabel": "no mês",
            "trend": trend(usd_delta),
            "tone": "negative" if usd_delta > 0 else "positive",
            "severity": "warning" if abs(usd_delta) > 2 else "normal",
            "icon": "currency",
            "source": "BCB SGS 3698",
            "updatedAt": as_of(usd_points),
        },
        {
            "id": "ipca",
            "label": "IPCA 12 MESES",
            "description": "Inflação acumulada em 12 meses",
            "value": f"{number(ipca)}%",
            "delta": signed(ipca_delta, 2, " p.p."),
            "deltaLabel": "no mês",
            "trend": trend(ipca_delta),
            "tone": "positive" if ipca_delta < 0 else "negative",
            "severity": "warning" if ipca > 4.5 else "normal",
            "icon": "inflation",
            "source": "BCB SGS 13522",
            "updatedAt": as_of(ipca12_points),
        },
        {
            "id": "selic",
            "label": "SELIC",
            "description": "Taxa básica de juros",
            "value": f"{number(selic)}%",
            "delta": signed(selic_delta, 2, " p.p."),
            "deltaLabel": "no mês",
            "trend": trend(selic_delta),
            "tone": "neutral" if selic_delta == 0 else "negative" if selic_delta > 0 else "positive",
            "severity": "warning" if selic >= 12 else "normal",
            "icon": "curve",
            "source": "BCB SGS 432",
            "updatedAt": as_of(selic_points),
        },
        {
            "id": "real-rate",
            "label": "JURO REAL EX-POST",
            "description": "Selic descontada da inflação",
            "value": f"{number(real_rate)}%",
            "delta": signed(real_delta, 2, " p.p."),
            "deltaLabel": "no mês",
            "trend": trend(real_delta),
            "tone": "negative" if real_delta > 0 else "positive",
            "severity": "warning" if real_rate > 7 else "normal",
            "icon": "curve",
            "source": "BCB SGS 432 + 13522",
            "updatedAt": as_of(real_points),
        },
        {
            "id": "ibov",
            "label": "IBOVESPA",
            "description": "Índice da bolsa brasileira",
            "value": f"{number(ibov, 0)} pts",
            "delta": signed(ibov_delta, 2, "%"),
            "deltaLabel": "no mês",
            "trend": trend(ibov_delta),
            "tone": "positive" if ibov_delta > 0 else "negative",
            "severity": "warning" if ibov_delta < -5 else "normal",
            "icon": "curve",
            "source": "Yahoo Finance",
            "updatedAt": as_of(ibov_points),
        },
        {
            "id": "ewz",
            "label": "EWZ",
            "description": "ETF Brasil em Nova York",
            "value": f"US$ {number(ewz)}",
            "delta": signed(ewz_delta, 2, "%"),
            "deltaLabel": "no mês",
            "trend": trend(ewz_delta),
            "tone": "positive" if ewz_delta > 0 else "negative",
            "severity": "warning" if ewz_delta < -5 else "normal",
            "icon": "curve",
            "source": "Yahoo Finance",
            "updatedAt": as_of(ewz_points),
        },
        {
            "id": "flow",
            "label": "FLUXO CAMBIAL",
            "description": "Conta financeira líquida",
            "value": f"-US$ {number(abs(flow), 1)} bi" if flow < 0 else f"US$ {number(flow, 1)} bi",
            "delta": "Saída" if flow_change < 0 else "Entrada",
            "deltaLabel": "no mês",
            "trend": "down" if flow_change < 0 else "up",
            "tone": "negative" if flow_change < 0 else "positive",
            "severity": "critical" if flow < -5 else "warning" if flow < 0 else "normal",
            "icon": "flow",
            "source": "BCB SGS 13961",
            "updatedAt": as_of(flow_points),
        },
        {
            "id": "janjometro",
            "label": "JANJÔMETRO",
            "description": f"Custo total publicado: {janjometro}"
            if janjometro
            else "Custo total publicado",
            "value": compact_brl(janjometro) if janjometro else "—",
            "delta": "Total",
            "deltaLabel": "publicado",
            "trend": "flat",
            "tone": "neutral",
            "severity": "normal",
            "icon": "currency",
            "source": "Janjômetro API",
            "updatedAt": NOW.isoformat().replace("+00:00", "Z"),
        },
        {
            "id": "interest",
            "label": "JUROS DA DÍVIDA PAGOS",
            "description": "Juros nominais em 12 meses",
            "value": f"R$ {number(interest_12m / 1000)} tri",
            "delta": signed(interest_delta, 1, "%"),
            "deltaLabel": "em 1 ano",
            "trend": trend(interest_delta),
            "tone": "negative" if interest_delta > 0 else "positive",
            "severity": "warning" if interest_12m >= 1000 else "normal",
            "icon": "risk",
            "source": "BCB SGS 4616",
            "updatedAt": as_of(interest_rolling),
        },
    ]

    flow_rolling_chart = chart_point_severity(
        flow_rolling,
        lambda value: "critical" if value < 0 else "normal",
    )
    debt_chart = chart_point_severity(
        debt_points[-24:],
        lambda value: "critical" if value > 80 else "warning" if value > 75 else "normal",
    )
    ipca_rolling_chart = chart_point_severity(
        ipca_rolling,
        lambda value: "critical" if value > 4.5 else "warning" if value > 4 else "normal",
    )
    gdp_growth_chart = chart_point_severity(
        gdp_growth_points,
        lambda value: "critical" if value < 0 else "normal",
    )
    primary_chart_points = chart_point_severity(
        primary_12m,
        lambda value: "critical" if value < 0 else "normal",
    )

    charts = [
        {
            "id": "usd-12m",
            "title": "DÓLAR PTAX — VELAS",
            "subtitle": "Velas mensais (OHLC) com média móvel de 12 meses",
            "kind": "candle",
            "section": "market",
            "series": [
                {
                    "id": "usd",
                    "label": "Dólar PTAX",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "prefix": "R$ ",
                    "points": usd_candles,
                }
            ],
        },
        {
            "id": "ewz-12m",
            "title": "EWZ — VELAS",
            "subtitle": "Velas mensais (OHLC) com média móvel de 12 meses",
            "kind": "candle",
            "section": "market",
            "series": [
                {
                    "id": "ewz",
                    "label": "EWZ",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "prefix": "US$ ",
                    "points": ewz_candles,
                }
            ],
        },
        {
            "id": "ibov-12m",
            "title": "IBOVESPA — VELAS",
            "subtitle": "Velas mensais (OHLC) com média móvel de 12 meses",
            "kind": "candle",
            "section": "market",
            "series": [
                {
                    "id": "ibov",
                    "label": "IBOVESPA",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 0,
                    "suffix": " pts",
                    "points": ibov_candles,
                }
            ],
        },
        {
            "id": "btc-usd-12m",
            "title": "BTC/USD — VELAS",
            "subtitle": "Velas mensais (OHLC) com média móvel de 12 meses",
            "kind": "candle",
            "section": "market",
            "series": [
                {
                    "id": "btc-usd",
                    "label": "BTC/USD",
                    "color": "#f5c542",
                    "axis": "left",
                    "decimals": 0,
                    "prefix": "US$ ",
                    "points": btc_usd_candles,
                }
            ],
        },
        {
            "id": "btc-brl-12m",
            "title": "BTC/BRL — VELAS",
            "subtitle": "Velas mensais (OHLC) com média móvel de 12 meses",
            "kind": "candle",
            "section": "market",
            "series": [
                {
                    "id": "btc-brl",
                    "label": "BTC/BRL",
                    "color": "#f5c542",
                    "axis": "left",
                    "decimals": 0,
                    "prefix": "R$ ",
                    "points": btc_brl_candles,
                }
            ],
        },
        {
            "id": "flow-12m",
            "title": "FLUXO CAMBIAL — ACUMULADO 12M",
            "subtitle": "Soma móvel de 12 meses (US$ bi)",
            "kind": "bar",
            "section": "market",
            "series": [
                {
                    "id": "flow",
                    "label": "Fluxo financeiro 12m",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 1,
                    "prefix": "US$ ",
                    "suffix": " bi",
                    "points": flow_rolling_chart,
                }
            ],
        },
        {
            "id": "debt-12m",
            "title": "DÍVIDA BRUTA DO GOVERNO GERAL",
            "subtitle": "Estoque em R$ tri (área) e em % do PIB (linha)",
            "kind": "line",
            "section": "fiscal",
            "series": [
                {
                    "id": "debt-brl",
                    "label": "Dívida bruta (R$ tri)",
                    "color": "#2de2e6",
                    "axis": "left",
                    "type": "area",
                    "decimals": 2,
                    "prefix": "R$ ",
                    "suffix": " tri",
                    "points": debt_brl_points[-24:],
                },
                {
                    "id": "debt",
                    "label": "Dívida bruta (% PIB)",
                    "color": "#f5c542",
                    "axis": "right",
                    "decimals": 2,
                    "suffix": "%",
                    "points": debt_chart,
                },
            ],
        },
        {
            "id": "revenue-12m",
            "title": "ARRECADAÇÃO FEDERAL — CRESCIMENTO ANUAL",
            "subtitle": "Acumulado em 12 meses vs. 1 ano antes (área), IPCA 12m e total em R$ tri",
            "kind": "line",
            "section": "fiscal",
            "series": [
                {
                    "id": "revenue-growth",
                    "label": "Crescimento anual",
                    "color": "#2de2e6",
                    "axis": "left",
                    "type": "area",
                    "decimals": 1,
                    "suffix": "%",
                    "points": revenue_growth,
                },
                {
                    "id": "revenue-ipca",
                    "label": "IPCA 12m",
                    "color": "#f5c542",
                    "axis": "left",
                    "dashed": True,
                    "decimals": 2,
                    "suffix": "%",
                    "points": ipca_matching(revenue_growth),
                },
                {
                    "id": "revenue",
                    "label": "Arrecadação 12m",
                    "color": "#c9d6d7",
                    "axis": "right",
                    "decimals": 2,
                    "prefix": "R$ ",
                    "suffix": " tri",
                    "points": to_trillions(revenue_rolling),
                },
            ],
        },
        {
            "id": "interest-12m",
            "title": "JUROS DA DÍVIDA — CRESCIMENTO ANUAL",
            "subtitle": "Juros nominais do setor público em 12 meses vs. 1 ano antes (área), IPCA 12m e total em R$ tri",
            "kind": "line",
            "section": "fiscal",
            "series": [
                {
                    "id": "interest-growth",
                    "label": "Crescimento anual",
                    "color": "#ff4d5e",
                    "axis": "left",
                    "type": "area",
                    "decimals": 1,
                    "suffix": "%",
                    "points": interest_growth,
                },
                {
                    "id": "interest-ipca",
                    "label": "IPCA 12m",
                    "color": "#f5c542",
                    "axis": "left",
                    "dashed": True,
                    "decimals": 2,
                    "suffix": "%",
                    "points": ipca_matching(interest_growth),
                },
                {
                    "id": "interest",
                    "label": "Juros 12m",
                    "color": "#c9d6d7",
                    "axis": "right",
                    "decimals": 2,
                    "prefix": "R$ ",
                    "suffix": " tri",
                    "points": to_trillions(interest_rolling),
                },
            ],
        },
        {
            "id": "interest-revenue",
            "title": "JUROS / ARRECADAÇÃO FEDERAL",
            "subtitle": "Juros nominais do governo central em 12m, em % da arrecadação federal em 12m",
            "kind": "line",
            "section": "fiscal",
            "series": [
                {
                    "id": "interest-share",
                    "label": "Juros / arrecadação",
                    "color": "#ff4d5e",
                    "axis": "left",
                    "type": "area",
                    "decimals": 1,
                    "suffix": "%",
                    "points": interest_share,
                }
            ],
        },
        {
            "id": "primary-12m",
            "title": "RESULTADO PRIMÁRIO E NOMINAL — ACUMULADO 12M",
            "subtitle": "Setor público consolidado, % do PIB (positivo = superávit); nominal inclui juros",
            "kind": "line",
            "section": "fiscal",
            "series": [
                {
                    "id": "primary",
                    "label": "Resultado primário 12m",
                    "color": "#ff4d5e",
                    "axis": "left",
                    "type": "area",
                    "decimals": 2,
                    "suffix": "% PIB",
                    "points": primary_chart_points,
                },
                {
                    "id": "nominal",
                    "label": "Resultado nominal 12m",
                    "color": "#f5c542",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "% PIB",
                    "points": nominal_12m,
                },
            ],
        },
        {
            "id": "gdp-growth",
            "title": "CRESCIMENTO DO PIB — ANUAL",
            "subtitle": "Variação real anual (%)",
            "kind": "bar",
            "section": "monetary",
            "series": [
                {
                    "id": "gdp-growth",
                    "label": "Crescimento do PIB",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": gdp_growth_chart,
                }
            ],
        },
        {
            "id": "ipca-12m",
            "title": "IPCA — ACUMULADO 12M",
            "subtitle": "Inflação composta em 12 meses (%)",
            "kind": "bar",
            "section": "monetary",
            "series": [
                {
                    "id": "ipca",
                    "label": "IPCA 12m",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": ipca_rolling_chart,
                }
            ],
        },
        {
            "id": "money-12m",
            "title": "IMPRESSÃO DE DINHEIRO — MÉDIA 12M",
            "subtitle": "Base monetária restrita, média móvel (R$ bi)",
            "kind": "line",
            "section": "monetary",
            "series": [
                {
                    "id": "money",
                    "label": "Base monetária (média 12m)",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 1,
                    "prefix": "R$ ",
                    "suffix": " bi",
                    "points": money_rolling,
                }
            ],
        },
        {
            "id": "real-rate-12m",
            "title": "SELIC x JURO REAL — MÉDIA 12M",
            "subtitle": "Média móvel de 12 meses (%)",
            "kind": "line",
            "section": "monetary",
            "series": [
                {
                    "id": "real-rate",
                    "label": "Juro real (média 12m)",
                    "color": "#ff4d5e",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": real_rolling,
                },
                {
                    "id": "selic",
                    "label": "Selic meta (média 12m)",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": selic_rolling,
                },
            ],
        },
        {
            "id": "pib-12m",
            "title": "PIB — ACUMULADO 12M",
            "subtitle": "Soma móvel de 12 meses (R$ trilhões)",
            "kind": "line",
            "section": "monetary",
            "series": [
                {
                    "id": "pib",
                    "label": "PIB 12m",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "prefix": "R$ ",
                    "suffix": " tri",
                    "points": pib_rolling,
                }
            ],
        },
        {
            "id": "focus-ipca",
            "title": "EXPECTATIVA DE INFLAÇÃO — ANUAL",
            "subtitle": "Mediana Focus do IPCA por ano (%)",
            "kind": "line",
            "section": "future",
            "series": [
                {
                    "id": "focus-ipca",
                    "label": "IPCA esperado",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": focus_ipca,
                }
            ],
        },
        {
            "id": "focus-pib",
            "title": "EXPECTATIVA DE PIB — ANUAL",
            "subtitle": "Mediana Focus de crescimento por ano (%)",
            "kind": "line",
            "section": "future",
            "series": [
                {
                    "id": "focus-pib",
                    "label": "PIB esperado",
                    "color": "#2de2e6",
                    "axis": "left",
                    "decimals": 2,
                    "suffix": "%",
                    "points": focus_pib,
                }
            ],
        },
        {
            "id": "di-curve",
            "title": "CURVA DE JUROS — DI FUTURO",
            "subtitle": "Taxa dos contratos DI de janeiro por vencimento (% a.a.)",
            "kind": "line",
            "section": "future",
            "series": di_curve,
        },
        {
            "id": "future-rates",
            "title": "JUROS FUTUROS — DI JAN/32 E JAN/35",
            "subtitle": "Taxa de fechamento por sessão (% a.a.)",
            "kind": "line",
            "section": "future",
            "series": di_daily,
        },
    ]

    alerts = build_alerts(
        debt,
        real_rate,
        flow,
        ipca,
        revenue,
        money_points,
        ewz_points,
        ibov_points,
        focus_ipca,
        daily=daily_market_alerts(di_sessions, ibov_sessions, usd_daily),
    )

    snapshot = {
        "generatedAt": NOW.isoformat().replace("+00:00", "Z"),
        "metrics": metrics,
        "charts": charts,
        "alerts": alerts,
        "news": news_items(),
        "impostometro": impostometro_live,
        "sources": STORE.sources,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    stale = [key for key, info in STORE.sources.items() if info["status"] != "ok"]

    print(f"Wrote {OUTPUT}")
    print(f"Sources: {len(STORE.sources) - len(stale)} ok, {len(stale)} not fresh")

    for key in stale:
        print(f"  {key}: {STORE.sources[key]['status']}")


if __name__ == "__main__":
    main()
